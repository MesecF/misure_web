"""Excel (.xlsx) and CSV file responses, shared by the measurement and article exports."""

import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATETIME_FORMAT = '%d.%m.%Y %H:%M:%S'


def local(moment):
    """Local time without time zone, as Excel expects."""
    return timezone.localtime(moment).replace(tzinfo=None) if moment else None


def is_formula_like(value):
    return isinstance(value, str) and value.startswith('=')


def xlsx_response(title, names, rows, filename):
    """
    rows: lists of typed values. Datetimes and decimals (measured values, one decimal place)
    get a number format; text, such as barcodes, stays text.
    Write-only workbook: rows are streamed, which keeps large exports fast. Column widths must
    be set before the first row, so the rows are collected first to measure them.
    """
    rows = list(rows)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title[:31])
    sheet.freeze_panes = 'A2'

    for index, name in enumerate(names):
        values = [row[index] for row in rows if row[index] is not None]
        if any(hasattr(value, 'hour') for value in values):
            width = 20
        else:
            longest = max([len(str(value)) for value in values], default=0)
            width = max(12, min(max(len(name), longest) + 2, 40))
        sheet.column_dimensions[get_column_letter(index + 1)].width = width

    def cell(value, **style):
        result = WriteOnlyCell(sheet, value)
        # openpyxl writes text starting with "=" as a formula; names and data are never formulas
        if is_formula_like(value):
            result.data_type = 's'
        for name, setting in style.items():
            setattr(result, name, setting)
        return result

    bold = Font(bold=True)
    sheet.append([cell(name, font=bold) for name in names])
    for row in rows:
        sheet.append([
            cell(value, number_format='dd.mm.yyyy hh:mm:ss') if hasattr(value, 'hour')
            else cell(value, number_format='0.0') if hasattr(value, 'as_tuple')
            else cell(value) if is_formula_like(value)
            else value
            for value in row
        ])

    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def csv_response(names, rows, filename):
    """';' delimiter, decimal comma and UTF-8 with BOM, so it opens correctly in Croatian Excel."""

    def text(value):
        if value is None:
            return ''
        if hasattr(value, 'hour'):
            return value.strftime(DATETIME_FORMAT)
        if hasattr(value, 'as_tuple'):
            return str(value).replace('.', ',')
        return str(value)

    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', lineterminator='\r\n')
    writer.writerow(names)
    for row in rows:
        writer.writerow([text(value) for value in row])
    response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response
