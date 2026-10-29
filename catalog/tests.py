from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Tenant, User
from measurements.models import Measurement, MeasurementValue

from .importers import import_articles, import_elements
from .models import Article, Category, CategoryElement, Element
from .validators import ean13_is_valid


class Ean13Tests(SimpleTestCase):
    def test_valid_barcodes(self):
        for code in ['4006381333931', '5901234123457', '9780201379624']:
            self.assertTrue(ean13_is_valid(code), code)

    def test_invalid_barcodes(self):
        for code in ['4006381333932', '400638133393', '40063813339310', '40063813339A1', '']:
            self.assertFalse(ean13_is_valid(code), code)


class CatalogTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant_a = Tenant.objects.create(code='A', name='Company A')
        cls.tenant_b = Tenant.objects.create(code='B', name='Company B')
        cls.admin_a = User.objects.create_user('admin', tenant=cls.tenant_a, role='admin')
        cls.user_a = User.objects.create_user('operator', tenant=cls.tenant_a, role='user')
        cls.admin_b = User.objects.create_user('admin', tenant=cls.tenant_b, role='admin')
        cls.category_a = Category.objects.create(tenant=cls.tenant_a, code='NIKE', name='Nike', created_by=cls.admin_a)
        cls.category_b = Category.objects.create(tenant=cls.tenant_b, code='NIKE', name='Nike B', created_by=cls.admin_b)

    def setUp(self):
        self.client.force_login(self.admin_a)


class CategoryViewTests(CatalogTestCase):
    def test_list_shows_only_own_categories(self):
        response = self.client.get(reverse('catalog:category_list'))
        self.assertEqual(list(response.context['object_list']), [self.category_a])

    def test_regular_user_is_denied(self):
        self.client.force_login(self.user_a)
        self.assertEqual(self.client.get(reverse('catalog:category_list')).status_code, 403)
        self.assertEqual(self.client.get(reverse('catalog:category_create')).status_code, 403)

    def test_create_sets_tenant_and_creator(self):
        response = self.client.post(reverse('catalog:category_create'), {'code': ' ADI ', 'name': 'Adidas'})
        self.assertRedirects(response, reverse('catalog:category_list'))
        category = Category.objects.get(tenant=self.tenant_a, code='ADI')
        self.assertEqual(category.created_by, self.admin_a)
        self.assertIsNone(category.updated_at)

    def test_duplicate_code_in_same_tenant_is_rejected(self):
        response = self.client.post(reverse('catalog:category_create'), {'code': 'NIKE', 'name': 'Again'})
        self.assertContains(response, 'A category with this code already exists.')

    def test_same_code_in_other_tenant_is_allowed(self):
        self.client.force_login(self.admin_b)
        response = self.client.post(reverse('catalog:category_create'), {'code': 'ADI', 'name': 'Adidas'})
        self.assertRedirects(response, reverse('catalog:category_list'))

    def test_update_records_who_changed(self):
        url = reverse('catalog:category_update', args=[self.category_a.pk])
        self.client.post(url, {'code': 'NIKE', 'name': 'Nike Inc.'})
        self.category_a.refresh_from_db()
        self.assertEqual(self.category_a.name, 'Nike Inc.')
        self.assertEqual(self.category_a.updated_by, self.admin_a)
        self.assertIsNotNone(self.category_a.updated_at)

    def test_saving_without_changes_does_not_record_update(self):
        url = reverse('catalog:category_update', args=[self.category_a.pk])
        self.client.post(url, {'code': 'NIKE', 'name': 'Nike'})
        self.category_a.refresh_from_db()
        self.assertIsNone(self.category_a.updated_at)

    def test_other_tenants_category_is_not_found(self):
        self.assertEqual(self.client.get(reverse('catalog:category_update', args=[self.category_b.pk])).status_code, 404)
        response = self.client.post(reverse('catalog:category_toggle_active', args=[self.category_b.pk]))
        self.assertEqual(response.status_code, 404)

    def test_toggle_active(self):
        url = reverse('catalog:category_toggle_active', args=[self.category_a.pk])
        self.client.post(url)
        self.category_a.refresh_from_db()
        self.assertFalse(self.category_a.is_active)
        self.client.post(url)
        self.category_a.refresh_from_db()
        self.assertTrue(self.category_a.is_active)

    def test_toggle_does_not_redirect_to_other_sites(self):
        url = reverse('catalog:category_toggle_active', args=[self.category_a.pk])
        response = self.client.post(url, {'next': 'https://example.com/'})
        self.assertRedirects(response, reverse('catalog:category_list'))


class ElementViewTests(CatalogTestCase):
    def test_create_element(self):
        data = {'code': 'L1', 'name': 'Length', 'unit': 'cm'}
        response = self.client.post(reverse('catalog:element_create'), data)
        self.assertRedirects(response, reverse('catalog:element_list'))
        element = Element.objects.get(code='L1')
        self.assertEqual((element.tenant, element.created_by), (self.tenant_a, self.admin_a))

    def test_duplicate_code_in_tenant_is_rejected_but_allowed_in_other_tenant(self):
        data = {'code': 'L1', 'name': 'Length', 'unit': 'cm'}
        self.client.post(reverse('catalog:element_create'), data)
        response = self.client.post(reverse('catalog:element_create'), {**data, 'name': 'Other'})
        self.assertContains(response, 'An element with this code already exists.')
        self.client.force_login(self.admin_b)
        self.assertRedirects(self.client.post(reverse('catalog:element_create'), data),
                             reverse('catalog:element_list'))

    def test_list_status_filter(self):
        common = {'tenant': self.tenant_a, 'unit': 'cm', 'created_by': self.admin_a}
        Element.objects.create(code='L1', name='Length', **common)
        Element.objects.create(code='W1', name='Width', is_active=False, **common)
        url = reverse('catalog:element_list')

        def codes(params):
            return [e.code for e in self.client.get(url, params).context['object_list']]

        self.assertEqual(codes({}), ['L1'])
        self.assertEqual(codes({'status': 'inactive'}), ['W1'])
        self.assertEqual(codes({'status': 'all'}), ['L1', 'W1'])

    def test_regular_user_is_denied(self):
        self.client.force_login(self.user_a)
        self.assertEqual(self.client.get(reverse('catalog:element_list')).status_code, 403)
        self.assertEqual(self.client.get(reverse('catalog:element_import')).status_code, 403)
        self.assertEqual(self.client.get(reverse('catalog:element_export', args=['csv'])).status_code, 403)

    def test_code_is_saved_in_upper_case_and_compared_regardless_of_case(self):
        self.client.post(reverse('catalog:element_create'), {'code': ' b05 ', 'name': 'Length', 'unit': 'cm'})
        self.assertEqual(Element.objects.get().code, 'B05')
        response = self.client.post(reverse('catalog:element_create'), {'code': 'B05', 'name': 'Other', 'unit': 'cm'})
        self.assertContains(response, 'An element with this code already exists.')

    def test_export_follows_status_filter_and_lists_categories(self):
        common = {'tenant': self.tenant_a, 'unit': 'cm', 'created_by': self.admin_a}
        length = Element.objects.create(code='L1', name='Length', **common)
        Element.objects.create(code='W1', name='Width', is_active=False, **common)
        other = Category.objects.create(tenant=self.tenant_a, code='ADIDAS', name='Adidas', created_by=self.admin_a)
        for category in (self.category_a, other):
            CategoryElement.objects.create(tenant=self.tenant_a, category=category, element=length, sort_order=10,
                                           created_by=self.admin_a)

        response = self.client.get(reverse('catalog:element_export', args=['xlsx']), {'status': 'all'})
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(rows[0][:5], ('Code', 'Name', 'Unit', 'Categories', 'Active'))
        self.assertEqual([row[:5] for row in rows[1:]], [
            ('L1', 'Length', 'cm', 'ADIDAS, NIKE', 'Yes'),
            ('W1', 'Width', 'cm', None, 'No'),
        ])

        lines = self.client.get(reverse('catalog:element_export', args=['csv'])).content.decode('utf-8-sig').splitlines()
        self.assertEqual(len(lines), 2)
        self.assertTrue(lines[1].startswith('L1;Length;cm;ADIDAS, NIKE;Yes;admin;'))


class ElementImportTests(CatalogTestCase):
    HEADER = 'code;name;unit\n'

    def run_import(self, rows, tenant=None, user=None):
        return import_elements((self.HEADER + rows).encode(), tenant or self.tenant_a, user or self.admin_a)

    def test_creates_elements_in_upper_case(self):
        result = self.run_import('b05;Lunghezza;cm\nC01;Vita;cm\n')
        self.assertEqual((result.created, result.rejected), (2, []))
        self.assertEqual(list(Element.objects.values_list('code', 'name', 'unit', 'created_by')),
                         [('B05', 'Lunghezza', 'cm', self.admin_a.pk), ('C01', 'Vita', 'cm', self.admin_a.pk)])

    def test_existing_codes_are_rejected_also_when_inactive(self):
        self.run_import('B05;Lunghezza;cm\nC01;Vita;cm\n')
        Element.objects.filter(code='C01').update(is_active=False)
        result = self.run_import('b05;Other;mm\nC01;Other;mm\nD;Nuovo;cm\n')
        self.assertEqual(result.created, 1)
        self.assertEqual([(row.line, row.reasons) for row in result.rejected], [
            (2, ['Element "B05" already exists.']),
            (3, ['Element "C01" already exists and is inactive.']),
        ])
        self.assertEqual(Element.objects.get(code='B05').name, 'Lunghezza')

    def test_invalid_rows_are_rejected(self):
        result = self.run_import(
            'A1;Name;cm\n'               # line 2: valid
            'A2;;cm\n'                   # 3: empty name
            'A3;Name\n'                  # 4: missing column
            'A4;Name;cm;x\n'             # 5: too many columns
            f'{"X" * 21};Name;cm\n'      # 6: code too long
            'a5;Name;cm\nA5;Other;cm\n'  # 7, 8: same code twice in the file
        )
        self.assertEqual(result.created, 1)
        self.assertEqual([(row.line, row.reasons) for row in result.rejected], [
            (3, ['Name is empty.']),
            (4, ['Expected 3 columns, found 2.']),
            (5, ['Too many columns.']),
            (6, ['Code is longer than 20 characters.']),
            (7, ['Code appears more than once in the file.']),
            (8, ['Code appears more than once in the file.']),
        ])

    def test_other_tenant_codes_do_not_matter(self):
        self.run_import('B05;Lunghezza;cm\n', self.tenant_b, self.admin_b)
        self.assertEqual(self.run_import('B05;Lunghezza;cm\n').created, 1)

    def test_upload_shows_result(self):
        file = SimpleUploadedFile('elements.csv', b'h\nB05;Lunghezza;cm\nB05;Again;cm\n')
        response = self.client.post(reverse('catalog:element_import'), {'file': file})
        self.assertEqual(response.context['result'].created, 0)
        self.assertContains(response, 'Code appears more than once in the file.')


class CategoryElementTests(CatalogTestCase):
    def setUp(self):
        super().setUp()
        common = {'tenant': self.tenant_a, 'unit': 'cm', 'created_by': self.admin_a}
        self.length = Element.objects.create(code='L', name='Length', **common)
        self.width = Element.objects.create(code='W', name='Width', **common)
        self.url = reverse('catalog:category_elements', args=[self.category_a.pk])

    def add(self, element, sort_order):
        return self.client.post(self.url, {'element': element.pk, 'sort_order': sort_order})

    def test_add_element_to_category(self):
        self.assertRedirects(self.add(self.length, 10), self.url)
        link = CategoryElement.objects.get()
        self.assertEqual((link.tenant, link.category, link.element, link.sort_order, link.created_by),
                         (self.tenant_a, self.category_a, self.length, 10, self.admin_a))

    def test_same_element_in_several_categories_with_own_order(self):
        category_2 = Category.objects.create(tenant=self.tenant_a, code='ADI', name='Adidas', created_by=self.admin_a)
        self.add(self.length, 10)
        self.client.post(reverse('catalog:category_elements', args=[category_2.pk]),
                         {'element': self.length.pk, 'sort_order': 50})
        self.assertEqual(sorted(CategoryElement.objects.values_list('category__code', 'sort_order')),
                         [('ADI', 50), ('NIKE', 10)])

    def test_suggested_sort_order(self):
        self.assertEqual(self.client.get(self.url).context['form']['sort_order'].value(), 10)
        self.add(self.length, 35)
        self.assertEqual(self.client.get(self.url).context['form']['sort_order'].value(), 45)

    def test_element_already_in_category_and_inactive_elements_are_not_offered(self):
        self.add(self.length, 10)
        Element.objects.create(tenant=self.tenant_a, code='X', name='Old', unit='cm', is_active=False,
                               created_by=self.admin_a)
        offered = list(self.client.get(self.url).context['form'].fields['element'].queryset)
        self.assertEqual(offered, [self.width])

    def test_other_tenants_element_cannot_be_added(self):
        other = Element.objects.create(tenant=self.tenant_b, code='L', name='Length', unit='cm',
                                       created_by=self.admin_b)
        response = self.add(other, 10)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(CategoryElement.objects.exists())

    def test_duplicate_sort_order_is_rejected(self):
        self.add(self.length, 10)
        response = self.add(self.width, 10)
        self.assertContains(response, 'This category already has an element with this sort order.')

    def test_change_sort_order(self):
        self.add(self.length, 10)
        self.add(self.width, 20)
        link = CategoryElement.objects.get(element=self.length)
        url = reverse('catalog:category_element_update', args=[link.pk])
        self.client.post(url, {'sort_order': 30})
        link.refresh_from_db()
        self.assertEqual((link.sort_order, link.updated_by), (30, self.admin_a))
        # A sort order taken by another element is refused
        self.client.post(url, {'sort_order': 20})
        link.refresh_from_db()
        self.assertEqual(link.sort_order, 30)

    def test_remove_element_from_category(self):
        self.add(self.length, 10)
        link = CategoryElement.objects.get()
        self.assertRedirects(self.client.post(reverse('catalog:category_element_delete', args=[link.pk])), self.url)
        self.assertFalse(CategoryElement.objects.exists())
        self.assertTrue(Element.objects.filter(pk=self.length.pk).exists())

    def test_other_tenants_category_is_not_found(self):
        self.assertEqual(
            self.client.get(reverse('catalog:category_elements', args=[self.category_b.pk])).status_code, 404)

    def test_active_elements_in_sort_order(self):
        self.add(self.width, 10)
        self.add(self.length, 20)
        self.assertEqual(CategoryElement.active_elements(self.category_a), [self.width, self.length])
        Element.objects.filter(pk=self.width.pk).update(is_active=False)
        self.assertEqual(CategoryElement.active_elements(self.category_a), [self.length])


class ArticleImportTests(CatalogTestCase):
    HEADER = 'code;color;size;category;barcode\n'

    def run_import(self, rows, encoding='utf-8'):
        return import_articles((self.HEADER + rows).encode(encoding), self.tenant_a, self.admin_a)

    def reasons(self, result):
        return {row.line: ' '.join(row.reasons) for row in result.rejected}

    def test_creates_articles_and_matches_category_case_insensitively(self):
        result = self.run_import('A1;BLACK;M;nike;4006381333931\nA2;WHITE;L;NIKE;5901234123457\n')
        self.assertEqual((result.created, result.updated, result.unchanged, result.rejected), (2, 0, 0, []))
        article = Article.objects.get(barcode='4006381333931')
        self.assertEqual((article.code, article.category, article.tenant, article.created_by),
                         ('A1', self.category_a, self.tenant_a, self.admin_a))

    def test_existing_barcode_is_updated_and_unchanged_is_counted(self):
        self.run_import('A1;BLACK;M;NIKE;4006381333931\nA2;WHITE;L;NIKE;5901234123457\n')
        result = self.run_import('A1;RED;M;NIKE;4006381333931\nA2;WHITE;L;NIKE;5901234123457\n')
        self.assertEqual((result.created, result.updated, result.unchanged), (0, 1, 1))
        article = Article.objects.get(barcode='4006381333931')
        self.assertEqual((article.color, article.updated_by), ('RED', self.admin_a))
        self.assertIsNone(Article.objects.get(barcode='5901234123457').updated_at)

    def test_inactive_article_is_updated_but_stays_inactive(self):
        self.run_import('A1;BLACK;M;NIKE;4006381333931\n')
        Article.objects.update(is_active=False)
        self.run_import('A1;RED;M;NIKE;4006381333931\n')
        article = Article.objects.get()
        self.assertEqual((article.color, article.is_active), ('RED', False))

    def test_invalid_rows_are_rejected_and_valid_ones_imported(self):
        Category.objects.create(tenant=self.tenant_a, code='OLD', name='Old', is_active=False, created_by=self.admin_a)
        result = self.run_import(
            'A1;BLACK;M;NIKE;4006381333931\n'  # line 2: valid
            'A2;BLACK;M;PUMA;5901234123457\n'  # 3: unknown category
            'A3;BLACK;M;OLD;5901234123457\n'  # 4: inactive category
            'A4;BLACK;M;NIKE;4006381333932\n'  # 5: wrong check digit
            'A5;;M;NIKE;9780201379624\n'  # 6: empty color
            'A6;BLACK;M;NIKE\n'  # 7: missing column
            'A7;VERYLONGCOLOR;M;NIKE;9780201379624\n'  # 8: color too long
            'A8;BLACK;M;NIKE;3,85011E+12\n'  # 9: barcode mangled by Excel
        )
        self.assertEqual(result.created, 1)
        reasons = self.reasons(result)
        self.assertEqual(sorted(reasons), [3, 4, 5, 6, 7, 8, 9])
        self.assertIn('does not exist', reasons[3])
        self.assertIn('inactive', reasons[4])
        self.assertIn('not a valid EAN-13', reasons[5])
        self.assertIn('Color is empty', reasons[6])
        self.assertIn('Expected 5 columns', reasons[7])
        self.assertIn('longer than 10', reasons[8])
        self.assertIn('not a valid EAN-13', reasons[9])

    def test_barcode_twice_in_file_rejects_both_rows(self):
        result = self.run_import('A1;BLACK;M;NIKE;4006381333931\nA2;WHITE;L;NIKE;4006381333931\n')
        self.assertEqual(result.created, 0)
        self.assertEqual(sorted(self.reasons(result)), [2, 3])

    def test_windows_1250_file_and_empty_rows(self):
        result = self.run_import('A1;ŽUTA;M;NIKE;4006381333931\n;;;;\n\n', encoding='cp1250')
        self.assertEqual((result.created, result.rejected), (1, []))
        self.assertEqual(Article.objects.get().color, 'ŽUTA')

    def test_trailing_empty_columns_are_allowed_but_extra_values_are_not(self):
        result = self.run_import('A1;BLACK;M;NIKE;4006381333931;;\nA2;BLACK;M;NIKE;5901234123457;x\n')
        self.assertEqual(result.created, 1)
        self.assertIn('Too many columns', self.reasons(result)[3])

    def test_code_color_and_size_are_converted_to_upper_case(self):
        self.run_import('ucc105h;black;xl;NIKE;4006381333931\n')
        article = Article.objects.get()
        self.assertEqual((article.code, article.color, article.size), ('UCC105H', 'BLACK', 'XL'))

    def test_same_variant_twice_in_file_rejects_both_rows(self):
        # Differs only in case, which does not matter after conversion to upper case
        result = self.run_import('A1;BLACK;M;NIKE;4006381333931\na1;black;m;NIKE;5901234123457\n')
        self.assertEqual(result.created, 0)
        reasons = self.reasons(result)
        self.assertEqual(sorted(reasons), [2, 3])
        self.assertIn('same code, color and size appear more than once', reasons[3])

    def test_variant_existing_under_other_barcode_is_rejected(self):
        self.run_import('A1;BLACK;M;NIKE;4006381333931\n')
        result = self.run_import('A1;BLACK;M;NIKE;5901234123457\n')
        self.assertEqual(result.created, 0)
        self.assertIn('already exists (barcode 4006381333931)', self.reasons(result)[2])

    def test_update_cannot_change_article_into_existing_variant(self):
        self.run_import('A1;BLACK;M;NIKE;4006381333931\nA2;BLACK;M;NIKE;5901234123457\n')
        result = self.run_import('A1;BLACK;M;NIKE;5901234123457\n')
        self.assertEqual(result.updated, 0)
        self.assertIn('already exists (barcode 4006381333931)', self.reasons(result)[2])
        self.assertEqual(Article.objects.get(barcode='5901234123457').code, 'A2')

    def test_import_does_not_touch_other_tenant(self):
        import_articles((self.HEADER + 'B1;BLACK;M;NIKE;4006381333931\n').encode(), self.tenant_b, self.admin_b)
        result = self.run_import('A1;BLACK;M;NIKE;4006381333931\n')
        self.assertEqual(result.created, 1)
        self.assertEqual(Article.objects.get(tenant=self.tenant_b).code, 'B1')


class InvalidImportFileTests(CatalogTestCase):
    def upload(self, url_name, content):
        file = SimpleUploadedFile('data.csv', content)
        return self.client.post(reverse(url_name), {'file': file})

    def test_non_csv_files_get_a_message_instead_of_an_error(self):
        files = [
            b'PK\x03\x04\x14\x00\x06\x00',                                 # Excel workbook (zip)
            'code;name;unit\nB05;Lunghezza;cm\n'.encode('utf-16'),        # UTF-16 text
            b'h\nA1;\x81\x98;M;NIKE;4006381333931\n',                      # neither UTF-8 nor Windows-1250
            b'h\n"' + b'x' * 140000 + b'\n',                               # quote never closed, huge field
        ]
        for url_name in ('catalog:article_import', 'catalog:element_import'):
            for content in files:
                response = self.upload(url_name, content)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, 'The file cannot be read as CSV.')
        self.assertFalse(Article.objects.exists() or Element.objects.exists())

    def test_reimport_of_unchanged_articles_runs_a_fixed_number_of_queries(self):
        def rows(count):
            barcodes = ['4006381333931', '5901234123457', '9780201379624', '4012345678901', '0012345678905']
            return ''.join(f'A{i};BLACK;M;NIKE;{barcodes[i]}\n' for i in range(count)).encode()

        import_articles(b'h\n' + rows(5), self.tenant_a, self.admin_a)
        with CaptureQueriesContext(connection) as two:
            import_articles(b'h\n' + rows(2), self.tenant_a, self.admin_a)
        with CaptureQueriesContext(connection) as five:
            result = import_articles(b'h\n' + rows(5), self.tenant_a, self.admin_a)
        self.assertEqual(result.unchanged, 5)
        self.assertEqual(len(two), len(five))


class ArticleViewTests(CatalogTestCase):
    def test_upload_shows_result(self):
        file = SimpleUploadedFile('articles.csv', b'h\nA1;BLACK;M;NIKE;4006381333931\nA2;X;M;PUMA;5901234123457\n')
        response = self.client.post(reverse('catalog:article_import'), {'file': file})
        self.assertEqual(response.context['result'].created, 1)
        self.assertContains(response, 'Category &quot;PUMA&quot; does not exist.')

    def test_regular_user_cannot_import(self):
        self.client.force_login(self.user_a)
        self.assertEqual(self.client.get(reverse('catalog:article_import')).status_code, 403)

    def test_manual_create_and_duplicate_barcode(self):
        data = {'code': 'A1', 'color': 'BLACK', 'size': 'M', 'category': self.category_a.pk, 'barcode': '4006381333931'}
        response = self.client.post(reverse('catalog:article_create'), data)
        self.assertRedirects(response, reverse('catalog:article_list'))
        response = self.client.post(reverse('catalog:article_create'), {**data, 'code': 'A2'})
        self.assertContains(response, 'An article with this barcode already exists.')

    def test_manual_create_converts_to_upper_case_and_rejects_existing_variant(self):
        data = {'code': 'a1', 'color': 'black', 'size': 'm', 'category': self.category_a.pk, 'barcode': '4006381333931'}
        self.client.post(reverse('catalog:article_create'), data)
        article = Article.objects.get()
        self.assertEqual((article.code, article.color, article.size), ('A1', 'BLACK', 'M'))
        response = self.client.post(reverse('catalog:article_create'), {**data, 'barcode': '5901234123457'})
        self.assertContains(response, 'An article with this code, color and size already exists (barcode 4006381333931).')

    def test_editing_article_keeps_its_own_variant(self):
        data = {'code': 'A1', 'color': 'BLACK', 'size': 'M', 'category': self.category_a.pk, 'barcode': '4006381333931'}
        self.client.post(reverse('catalog:article_create'), data)
        url = reverse('catalog:article_update', args=[Article.objects.get().pk])
        self.assertRedirects(self.client.post(url, data), reverse('catalog:article_list'))

    def test_manual_create_rejects_invalid_barcode(self):
        data = {'code': 'A1', 'color': 'BLACK', 'size': 'M', 'category': self.category_a.pk, 'barcode': '4006381333932'}
        response = self.client.post(reverse('catalog:article_create'), data)
        self.assertContains(response, 'is not a valid EAN-13 barcode.')

    def test_list_search_and_filters(self):
        import_articles(b'h\nA1;BLACK;M;NIKE;4006381333931\nB2;BLACK;M;NIKE;5901234123457\n',
                        self.tenant_a, self.admin_a)
        Article.objects.filter(code='B2').update(is_active=False)
        url = reverse('catalog:article_list')

        def codes(params):
            return [a.code for a in self.client.get(url, params).context['object_list']]

        self.assertEqual(codes({}), ['A1'])
        self.assertEqual(codes({'status': 'all'}), ['A1', 'B2'])
        self.assertEqual(codes({'status': 'all', 'q': '590123'}), ['B2'])

    def test_category_code_is_unique_regardless_of_case(self):
        response = self.client.post(reverse('catalog:category_create'), {'code': 'nike', 'name': 'Lowercase'})
        self.assertContains(response, 'A category with this code already exists.')

    def test_export_follows_list_filters(self):
        import_articles(b'h\nA1;BLACK;M;NIKE;4006381333931\nB2;BLACK;M;NIKE;5901234123457\n',
                        self.tenant_a, self.admin_a)
        import_articles(b'h\nC3;RED;S;NIKE;9780201379624\n', self.tenant_b, self.admin_b)
        Article.objects.filter(code='B2').update(is_active=False)

        response = self.client.get(reverse('catalog:article_export', args=['xlsx']), {'status': 'all'})
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(rows[0][:7], ('Code', 'Color', 'Size', 'Category', 'Category name', 'Barcode', 'Active'))
        self.assertEqual([row[:7] for row in rows[1:]], [
            ('A1', 'BLACK', 'M', 'NIKE', 'Nike', '4006381333931', 'Yes'),
            ('B2', 'BLACK', 'M', 'NIKE', 'Nike', '5901234123457', 'No'),
        ])
        self.assertEqual(rows[1][7], 'admin')

        response = self.client.get(reverse('catalog:article_export', args=['csv']))
        lines = response.content.decode('utf-8-sig').strip().splitlines()
        self.assertEqual(len(lines), 2)
        self.assertTrue(lines[1].startswith('A1;BLACK;M;NIKE;Nike;4006381333931;Yes;admin;'))

    def test_regular_user_cannot_export(self):
        self.client.force_login(self.user_a)
        self.assertEqual(self.client.get(reverse('catalog:article_export', args=['csv'])).status_code, 403)
        self.client.force_login(self.admin_a)
        self.assertEqual(self.client.get(reverse('catalog:article_export', args=['pdf'])).status_code, 404)


class UsedRecordLockTests(CatalogTestCase):
    """Code and unit of a measured element and the code of a measured category cannot be changed."""

    def setUp(self):
        super().setUp()
        self.element = Element.objects.create(tenant=self.tenant_a, code='L1', name='Length', unit='cm',
                                              created_by=self.admin_a)
        self.article = Article.objects.create(tenant=self.tenant_a, code='A1', color='BLACK', size='M',
                                              category=self.category_a, barcode='4006381333931',
                                              created_by=self.admin_a)

    def measure(self):
        measurement = Measurement(tenant=self.tenant_a, article=self.article, employee_number=1,
                                  created_by=self.admin_a)
        measurement.copy_article_data()
        measurement.save()
        value = MeasurementValue(measurement=measurement, element=self.element, value='12.5')
        value.copy_element_data()
        value.save()

    def test_unused_element_can_be_changed(self):
        url = reverse('catalog:element_update', args=[self.element.pk])
        self.client.post(url, {'code': 'L2', 'name': 'Length 2', 'unit': 'mm'})
        self.element.refresh_from_db()
        self.assertEqual((self.element.code, self.element.name, self.element.unit), ('L2', 'Length 2', 'mm'))

    def test_measured_element_keeps_code_and_unit(self):
        self.measure()
        url = reverse('catalog:element_update', args=[self.element.pk])
        response = self.client.get(url)
        self.assertContains(response, 'Cannot be changed because the element has measured values.', count=2)
        self.assertTrue(response.context['form'].fields['code'].disabled)
        # Posted code and unit are ignored, the name is saved
        self.assertRedirects(self.client.post(url, {'code': 'L2', 'name': 'Total length', 'unit': 'mm'}),
                             reverse('catalog:element_list'))
        self.element.refresh_from_db()
        self.assertEqual((self.element.code, self.element.name, self.element.unit), ('L1', 'Total length', 'cm'))

    def test_measured_category_keeps_code(self):
        url = reverse('catalog:category_update', args=[self.category_a.pk])
        self.assertFalse(self.client.get(url).context['form'].fields['code'].disabled)
        self.measure()
        response = self.client.get(url)
        self.assertContains(response, 'The code cannot be changed because the category has measurements.')
        self.client.post(url, {'code': 'ADIDAS', 'name': 'Nike shoes'})
        self.category_a.refresh_from_db()
        self.assertEqual((self.category_a.code, self.category_a.name), ('NIKE', 'Nike shoes'))
