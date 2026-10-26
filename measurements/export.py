"""Export of the filtered measurements to Excel (.xlsx) and CSV."""

import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DATETIME_FORMAT = '%d.%m.%Y %H:%M:%S'


def header(columns, include_deleted):
    names = [_('Measured at'), _('Employee number'), _('Article code'), _('Color'), _('Size'), _('Barcode'),
             _('Category')]
    names += [f'{element.name} ({element.unit})' for element in columns]
    names += [_('Entered by'), _('Edited by'), _('Edited at')]
    if include_deleted:
        names += [_('Deleted by'), _('Deleted at')]
    return names


def local(moment):
    """Local time without time zone, as Excel expects."""
    return timezone.localtime(moment).replace(tzinfo=None) if moment else None


def rows(measurements, include_deleted):
    """Rows with typed values: datetimes, numbers and text."""
    for m in measurements:
        row = [local(m.measured_at), m.employee_number, m.article_code, m.article_color, m.article_size,
               m.article_barcode, m.category_code]
        row += m.element_values
        row += [m.created_by.username, m.updated_by.username if m.updated_by else None, local(m.updated_at)]
        if include_deleted:
            row += [m.deleted_by.username if m.deleted_by else None, local(m.deleted_at)]
        yield row


def to_xlsx(measurements, columns, include_deleted, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _('Measurements')[:31]
    names = header(columns, include_deleted)
    sheet.append(names)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'

    first_value_column = 8
    value_columns = range(first_value_column, first_value_column + len(columns))
    for row in rows(measurements, include_deleted):
        sheet.append(row)
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(row=sheet.max_row, column=index)
            if index in value_columns:
                cell.number_format = '0.0'
            elif hasattr(value, 'hour'):
                cell.number_format = 'dd.mm.yyyy hh:mm:ss'

    for index, name in enumerate(names, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(len(name) + 2, 40))
    sheet.column_dimensions['A'].width = 20

    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def to_csv(measurements, columns, include_deleted, filename):
    """';' delimiter, decimal comma and UTF-8 with BOM, so it opens correctly in Croatian Excel."""

    def text(value):
        if value is None:
            return ''
        if hasattr(value, 'hour'):
            return value.strftime(DATETIME_FORMAT)
        if hasattr(value, 'as_tuple'):
            return str(value).replace('.', ',')
        return str(value)

    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', lineterminator='\r\n')
    writer.writerow(header(columns, include_deleted))
    for row in rows(measurements, include_deleted):
        writer.writerow([text(value) for value in row])
    response = HttpResponse(output.getvalue().encode('utf-8-sig'), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response
